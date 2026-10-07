"""
Resume Parser Module
====================
Handles text extraction from PDF and DOCX resume files.
Uses pdfplumber for PDF files and python-docx for DOCX files.
"""

import pdfplumber
from docx import Document
import io


def extract_text_from_pdf(file) -> str:
    """
    Extract text content from a PDF file.
    
    Args:
        file: Uploaded file object (BytesIO or file-like object)
    
    Returns:
        str: Extracted text from all pages of the PDF
    """
    text = ""
    try:
        with pdfplumber.open(file) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        raise ValueError(f"Error reading PDF file: {str(e)}")
    
    if not text.strip():
        raise ValueError("Could not extract any text from the PDF. The file may be image-based or corrupted.")
    
    return text.strip()


def extract_text_from_docx(file) -> str:
    """
    Extract text content from a DOCX file.
    
    Args:
        file: Uploaded file object (BytesIO or file-like object)
    
    Returns:
        str: Extracted text from all paragraphs and tables in the DOCX
    """
    text = ""
    try:
        doc = Document(file)
        
        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"
        
        # Also extract text from tables (resumes often use tables for layout)
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text += " | ".join(row_text) + "\n"
                    
    except Exception as e:
        raise ValueError(f"Error reading DOCX file: {str(e)}")
    
    if not text.strip():
        raise ValueError("Could not extract any text from the DOCX file.")
    
    return text.strip()


def extract_text(file, filename: str) -> str:
    """
    Main function to extract text from a resume file.
    Determines file type based on extension and calls the appropriate extractor.
    
    Args:
        file: Uploaded file object
        filename: Name of the uploaded file
    
    Returns:
        str: Extracted text from the resume
    """
    filename_lower = filename.lower()
    
    if filename_lower.endswith('.pdf'):
        return extract_text_from_pdf(file)
    elif filename_lower.endswith('.docx'):
        return extract_text_from_docx(file)
    else:
        raise ValueError("Unsupported file format. Please upload a PDF or DOCX file.")
